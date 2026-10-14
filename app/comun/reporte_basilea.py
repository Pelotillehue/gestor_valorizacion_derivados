"""!
@package comun.reporte_basilea
@brief Reporte FX Forward Basilea: detalle de forwards vigentes, cuadros resumen y uso de líneas.

Sigue la planilla "Operaciones Forward Basilea": la hoja Basilea con una fila por forward y su
valor razonable activo y pasivo, la hoja Resumen con tablas dinámicas sobre ese detalle y la
hoja Líneas con la situación de cada contraparte. Los montos de la moneda principal (USD o UF)
y de la secundaria no se suman entre pares distintos para no mezclar monedas.
"""

import io
from datetime import date, datetime
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import TABLOID, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from comun.formato import MODALIDADES, numero, rut
from comun.reporte_pnl import AZUL, DENTRO, FORMATO_MONTO, LLAMADO, ROJO

## Días al vencimiento hasta los que un forward se considera de corto plazo ("Hasta 7 Días").
DIAS_CORTO_PLAZO = 7
## Columnas de montos de la hoja Basilea.
MONTOS_BASILEA = [
    "Monto Moneda Principal", "Monto Moneda Secundaria", "Valor Razonable Activo",
    "Valor Razonable Pasivo", "Valor Razonable",
]  # fmt: skip


def _moneda(codigo: str) -> str:
    """!
    @brief Nombre visible de una moneda: la CLF se muestra como UF.
    @param codigo Código de moneda.
    @return Nombre visible.
    """
    return "UF" if codigo == "CLF" else codigo


def detalle(operaciones: list[dict]) -> pd.DataFrame:
    """!
    @brief Hoja Basilea: una fila por FX Forward vigente.
    @param operaciones Operaciones de `GET /cartera/forwards-basilea`.
    @return DataFrame con las columnas de la hoja Basilea.
    """
    filas = []
    for o in operaciones:
        principal, secundaria = _moneda(o["moneda_principal"]), _moneda(o["moneda_secundaria"])
        vr = o["valor_razonable"]
        filas.append(
            {
                "Movimiento": o["movimiento"],
                "Tipo de Forward": f"FWD {principal}/{secundaria}",
                "Modalidad": MODALIDADES.get(o["modalidad_liquidacion"], o["modalidad_liquidacion"]),
                "Folio": o["deal_number"],
                "Rut Contraparte": rut(o.get("rut_contraparte")),
                "Contraparte": o.get("contraparte") or "",
                "Calce": "Calzada" if o.get("grupo_calce") else "Descalzada",
                "Moneda Principal": principal,
                "Monto Moneda Principal": o["monto_principal"],
                "Tipo de Cambio": o["tipo_cambio"],
                "Moneda Secundaria": secundaria,
                "Monto Moneda Secundaria": o["monto_secundario"],
                "Fecha Inicio": pd.Timestamp(o["fecha_negociacion"]),
                "Plazo": o["plazo"],
                "Fecha Vencimiento": pd.Timestamp(o["fecha_vencimiento"]),
                "Días al Vencimiento": o["dias_al_vencimiento"],
                "Tasa VR Activo (%)": None if o["tasa_activo"] is None else o["tasa_activo"] * 100,
                "Valor Razonable Activo": o["valor_razonable_activo"],
                "Tasa VR Pasivo (%)": None if o["tasa_pasivo"] is None else o["tasa_pasivo"] * 100,
                "Valor Razonable Pasivo": o["valor_razonable_pasivo"],
                "Valor Razonable": vr,
                "Utilidad / Pérdida": "" if vr is None else ("Utilidad" if vr > 0 else "Pérdida"),
                "Hasta 7 Días": "SI" if o["dias_al_vencimiento"] <= DIAS_CORTO_PLAZO else "NO",
                "Error": o.get("error") or "",
            }
        )
    return pd.DataFrame(filas)


def _por_movimiento(datos: pd.DataFrame, filas: list[str], valor: str, descalce: bool = False) -> pd.DataFrame:
    """!
    @brief Tabla dinámica con columnas Compra y Venta.
    @param datos Detalle Basilea.
    @param filas Columnas que forman las filas.
    @param valor Columna que se suma.
    @param descalce Si es verdadero, agrega Descalce = Compra − Venta; si no, Total = Compra + Venta.
    @return DataFrame con las filas como columnas comunes.
    """
    tabla = datos.pivot_table(index=filas, columns="Movimiento", values=valor, aggfunc="sum", fill_value=0.0)
    tabla = tabla.reindex(columns=["Compra", "Venta"], fill_value=0.0)
    tabla.columns.name = None
    if descalce:
        tabla["Descalce"] = tabla["Compra"] - tabla["Venta"]
    else:
        tabla["Total"] = tabla["Compra"] + tabla["Venta"]
    return tabla.reset_index()


def _con_total(tabla: pd.DataFrame, etiqueta_col: str, etiqueta: str = "Total general") -> pd.DataFrame:
    """!
    @brief Agrega una fila de totales.
    @param tabla Tabla con una columna de etiquetas y columnas numéricas.
    @param etiqueta_col Columna donde va la etiqueta del total.
    @param etiqueta Texto de la fila.
    @return La tabla con la fila agregada.
    """
    numericas = tabla.select_dtypes("number").columns
    total = {c: "" for c in tabla.columns} | {etiqueta_col: etiqueta} | {c: tabla[c].sum() for c in numericas}
    return pd.concat([tabla, pd.DataFrame([total])], ignore_index=True)


def _una_moneda(datos: pd.DataFrame, columna: str) -> bool:
    """!
    @brief Verdadero si todas las filas están en la misma moneda, de modo que se pueden sumar.
    @param datos Detalle Basilea.
    @param columna Columna de moneda.
    @return Verdadero si hay una sola moneda.
    """
    return datos[columna].nunique() <= 1


def resumen(datos: pd.DataFrame) -> list[tuple[str, pd.DataFrame]]:
    """!
    @brief Hoja Resumen: las tablas dinámicas de la planilla, sobre los forwards valorizados.

    Los montos en moneda principal o secundaria se agrupan siempre por tipo de forward y sólo
    llevan total general si todos los tipos comparten moneda.
    @param datos Resultado de @ref comun.reporte_basilea.detalle.
    @return Lista de pares (título, tabla).
    """
    valorizadas = datos[datos["Valor Razonable"].notna()]
    tipo = "Tipo de Forward"
    cuadros = []

    vr_tipo = _con_total(_por_movimiento(valorizadas, [tipo], "Valor Razonable"), tipo)
    cuadros.append(("Valor razonable (CLP) por tipo de forward", vr_tipo))

    secundaria = _por_movimiento(datos, [tipo, "Moneda Secundaria"], "Monto Moneda Secundaria")
    if _una_moneda(datos, "Moneda Secundaria"):
        secundaria = _con_total(secundaria, tipo)
    cuadros.append(("Monto en moneda secundaria por tipo de forward", secundaria))

    principal = []
    for (nombre, mon), grupo in datos.groupby([tipo, "Moneda Principal"], sort=True):
        por_fecha = _por_movimiento(grupo, ["Fecha Vencimiento"], "Monto Moneda Principal", descalce=True)
        por_fecha.insert(0, tipo, "")
        por_fecha.insert(1, "Moneda Principal", mon)
        subtotal = {tipo: nombre, "Moneda Principal": mon, "Fecha Vencimiento": pd.NaT}
        subtotal |= {c: por_fecha[c].sum() for c in ("Compra", "Venta", "Descalce")}
        principal += [pd.DataFrame([subtotal]), por_fecha]
    principal = pd.concat(principal, ignore_index=True) if principal else pd.DataFrame()
    cuadros.append(("Monto en moneda principal por tipo de forward y vencimiento, con descalce", principal))

    plazo = valorizadas.pivot_table(
        index="Hasta 7 Días", columns="Utilidad / Pérdida", values="Valor Razonable", aggfunc="sum", fill_value=0.0
    ).reindex(index=["NO", "SI"], columns=["Pérdida", "Utilidad"], fill_value=0.0)
    plazo.columns.name = None
    plazo["Total"] = plazo.sum(axis=1)
    cuadros.append(("Valor razonable (CLP) por plazo al vencimiento", _con_total(plazo.reset_index(), "Hasta 7 Días")))

    resultado = valorizadas.pivot_table(
        index="Utilidad / Pérdida", columns=tipo, values="Valor Razonable", aggfunc="sum", fill_value=0.0
    ).reindex(index=["Pérdida", "Utilidad"], fill_value=0.0)
    resultado["Total"] = resultado.sum(axis=1)
    resultado.columns.name = None
    cuadros.append(
        ("Valor razonable (CLP) por resultado y tipo de forward", _con_total(resultado.reset_index(), "Utilidad / Pérdida"))
    )

    modalidad_p = _por_movimiento(datos, [tipo, "Moneda Principal", "Modalidad"], "Monto Moneda Principal", descalce=True)
    if _una_moneda(datos, "Moneda Principal"):
        modalidad_p = _con_total(modalidad_p, tipo)
    cuadros.append(("Monto en moneda principal por modalidad, con descalce", modalidad_p))

    modalidad_s = _por_movimiento(datos, [tipo, "Moneda Secundaria", "Modalidad"], "Monto Moneda Secundaria")
    if _una_moneda(datos, "Moneda Secundaria"):
        modalidad_s = _con_total(modalidad_s, tipo)
    cuadros.append(("Monto en moneda secundaria por modalidad", modalidad_s))
    return cuadros


def lineas(contrapartes: list[dict], garantias: list[dict], datos: pd.DataFrame) -> pd.DataFrame:
    """!
    @brief Hoja Líneas: una fila por contraparte con su línea, su posición en forwards y su situación.

    La exposición y la situación de garantías consideran todas las operaciones vigentes de la
    contraparte, igual que el control de garantías; el valor razonable y las posiciones netas,
    sólo sus FX Forward.
    @param contrapartes Respuesta de `GET /contrapartes`.
    @param garantias Contrapartes de `GET /cartera/garantias`.
    @param datos Resultado de @ref comun.reporte_basilea.detalle.
    @return DataFrame ordenado por contraparte.
    """
    por_rut = {g["rut"]: g for g in garantias}
    datos = datos.assign(
        _rut=datos["Rut Contraparte"],
        _neto=datos["Monto Moneda Principal"].where(datos["Movimiento"] == "Compra", -datos["Monto Moneda Principal"]),
    )
    monedas = sorted(datos["Moneda Principal"].unique())
    filas = []
    for c in contrapartes:
        propias = datos[datos["_rut"] == rut(c["rut"])]
        g = por_rut.get(c["rut"], {})
        exposicion = g.get("exposicion", 0.0)
        fila = {
            "Rut Contraparte": rut(c["rut"]),
            "Contraparte": c["nombre"],
            "Forwards Vigentes": len(propias),
            **{f"Posición Neta {m}": propias.loc[propias["Moneda Principal"] == m, "_neto"].sum() for m in monedas},
            "Valor Razonable FX Fwd (CLP)": propias["Valor Razonable"].sum(min_count=1),
            "Exposición Total (CLP)": exposicion,
            "Monto Línea (CLP)": c["monto_linea"],
            "Uso de Línea (%)": (max(exposicion, 0.0) / c["monto_linea"] * 100) if c["monto_linea"] else None,
            "Línea Disponible (CLP)": c["monto_linea"] - max(exposicion, 0.0),
            "MTA (CLP)": c["mta"],
            "Situación": LLAMADO if g.get("llamado") else DENTRO,
            "Monto a Llamar (CLP)": g.get("monto_llamado", 0.0),
        }
        filas.append(fila)
    tabla = pd.DataFrame(filas)
    if tabla.empty:
        return tabla
    tabla["Valor Razonable FX Fwd (CLP)"] = tabla["Valor Razonable FX Fwd (CLP)"].fillna(0.0)
    return tabla.sort_values("Contraparte").reset_index(drop=True)


def _encabezado(fila) -> None:
    """!
    @brief Aplica el estilo de encabezado (azul, texto blanco) a una fila de celdas.
    @param fila Celdas.
    """
    borde = Side(style="thin", color="BFBFBF")
    for celda in fila:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = Border(left=borde, right=borde, top=borde, bottom=borde)


def _formatear(hoja, tabla: pd.DataFrame, fila_inicio: int) -> None:
    """!
    @brief Formatos de número y fecha de una tabla escrita en una hoja.
    @param hoja Hoja de openpyxl.
    @param tabla Tabla escrita.
    @param fila_inicio Fila (1-based) del encabezado.
    """
    for i, columna in enumerate(tabla.columns, start=1):
        letra = get_column_letter(i)
        if columna in ("Tipo de Cambio",) or "(%)" in columna:
            formato = "#,##0.0000"
        elif pd.api.types.is_datetime64_any_dtype(tabla[columna]):
            formato = "dd-mm-yyyy"
        elif pd.api.types.is_numeric_dtype(tabla[columna]) and columna not in ("Plazo", "Días al Vencimiento", "Forwards Vigentes"):
            formato = FORMATO_MONTO
        else:
            continue
        for f in range(fila_inicio + 1, fila_inicio + len(tabla) + 1):
            hoja[f"{letra}{f}"].number_format = formato


def excel(datos: pd.DataFrame, cuadros: list[tuple[str, pd.DataFrame]], tabla_lineas: pd.DataFrame, fecha: date) -> bytes:
    """!
    @brief Excel con las hojas Resumen, Basilea y Líneas.
    @param datos Resultado de @ref comun.reporte_basilea.detalle.
    @param cuadros Resultado de @ref comun.reporte_basilea.resumen.
    @param tabla_lineas Resultado de @ref comun.reporte_basilea.lineas.
    @param fecha Fecha de proceso.
    @return Contenido del archivo .xlsx.
    """
    salida = io.BytesIO()
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        libro = escritor.book
        hoja = libro.create_sheet("Resumen")
        hoja["B2"] = "Resumen FX Forward Basilea"
        hoja["B2"].font = Font(bold=True, size=14)
        hoja["B3"] = f"Fecha de proceso {fecha:%d-%m-%Y}. Valor razonable en CLP, desde la institución."
        hoja["B3"].font = Font(color="52514E")
        fila = 5
        for titulo, tabla in cuadros:
            hoja.cell(fila, 2, titulo).font = Font(bold=True)
            tabla = tabla.copy()
            for c in tabla.columns:
                if pd.api.types.is_datetime64_any_dtype(tabla[c]):
                    tabla[c] = tabla[c].dt.date
            tabla.to_excel(escritor, sheet_name="Resumen", index=False, startrow=fila, startcol=1)
            _encabezado(hoja[fila + 1][1 : 1 + len(tabla.columns)])
            for i, columna in enumerate(tabla.columns, start=2):
                letra = get_column_letter(i)
                for f in range(fila + 2, fila + 2 + len(tabla)):
                    celda = hoja[f"{letra}{f}"]
                    if columna == "Fecha Vencimiento":
                        celda.number_format = "dd-mm-yyyy"
                    elif isinstance(celda.value, (int, float)):
                        celda.number_format = FORMATO_MONTO
            for k, fila_tabla in enumerate(tabla.itertuples(index=False)):
                # Totales y subtotales (en el cuadro por vencimiento, la fila del tipo de forward) en negrita.
                subtotal = "Fecha Vencimiento" in tabla.columns and pd.isna(fila_tabla[tabla.columns.get_loc("Fecha Vencimiento")])
                if fila_tabla[0] == "Total general" or subtotal:
                    for celda in hoja[fila + 2 + k][1 : 1 + len(tabla.columns)]:
                        celda.font = Font(bold=True)
            fila += len(tabla) + 4
        for i in range(2, 10):
            hoja.column_dimensions[get_column_letter(i)].width = 22

        basilea = datos.drop(columns=["Error"]).copy()
        for nombre, tabla in (("Basilea", basilea), ("Líneas", tabla_lineas)):
            tabla.to_excel(escritor, sheet_name=nombre, index=False, startrow=3)
            h = escritor.sheets[nombre]
            h["A1"] = "Detalle Basilea" if nombre == "Basilea" else "Líneas por contraparte"
            h["A1"].font = Font(bold=True, size=14)
            h["A2"] = f"Fecha de proceso {fecha:%d-%m-%Y}. Montos en CLP salvo los de moneda principal y secundaria."
            h["A2"].font = Font(color="52514E")
            _encabezado(h[4])
            _formatear(h, tabla, 4)
            for i, columna in enumerate(tabla.columns, start=1):
                h.column_dimensions[get_column_letter(i)].width = 26 if columna == "Contraparte" else 16
            h.freeze_panes = "E5" if nombre == "Basilea" else "C5"
            if nombre == "Líneas" and "Situación" in tabla.columns:
                col = get_column_letter(list(tabla.columns).index("Situación") + 1)
                for f in range(5, len(tabla) + 5):
                    if h[f"{col}{f}"].value == LLAMADO:
                        h[f"{col}{f}"].font = Font(bold=True, color=ROJO)
        libro.move_sheet("Resumen", offset=-libro.index(libro["Resumen"]))
    return salida.getvalue()


## Columnas que se muestran con cuatro decimales.
_CUATRO_DECIMALES = ("Tipo de Cambio", "Tasa VR Activo (%)", "Tasa VR Pasivo (%)")
## Columnas enteras que no son montos.
_ENTEROS = ("Plazo", "Días al Vencimiento", "Forwards Vigentes")


def _texto(columna: str, valor) -> str:
    """!
    @brief Valor de una celda como texto para el PDF, con formato chileno.
    @param columna Nombre de la columna.
    @param valor Valor.
    @return Texto.
    """
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    if isinstance(valor, (pd.Timestamp, datetime, date)):
        return f"{valor:%d-%m-%Y}"
    if isinstance(valor, (int, float)):
        if columna in _CUATRO_DECIMALES:
            return numero(valor, 4)
        if columna == "Uso de Línea (%)":
            return numero(valor, 1)
        return numero(valor, 0)
    return str(valor)


def _tabla_pdf(tabla: pd.DataFrame, ancho: float, tamano: float = 7) -> LongTable:
    """!
    @brief Tabla de reportlab con encabezado azul, montos a la derecha, negativos en rojo y totales en negrita.

    El ancho de cada columna es proporcional al texto más largo que contiene, con un mínimo y
    un máximo, y el conjunto ocupa el ancho indicado.
    @param tabla Tabla a dibujar.
    @param ancho Ancho total disponible, en puntos.
    @param tamano Tamaño de letra del cuerpo.
    @return `LongTable` que repite el encabezado en cada página.
    """
    encabezado = ParagraphStyle(
        "enc", fontName="Helvetica-Bold", fontSize=tamano, leading=tamano + 1.5, textColor=colors.white, alignment=1
    )
    celdas = [[_texto(c, v) for c, v in zip(tabla.columns, fila)] for fila in tabla.itertuples(index=False)]
    largos = []
    for i, columna in enumerate(tabla.columns):
        contenido = max([len(f[i]) for f in celdas] or [0])
        # Al menos la palabra más larga del encabezado, para no cortarla.
        palabra = max(len(p) for p in columna.split()) + 1
        largos.append(min(max(contenido, min(len(columna), 12), palabra, 4), 26))
    total = sum(largos)
    anchos = [ancho * l / total for l in largos]

    datos = [[Paragraph(escape(c), encabezado) for c in tabla.columns], *celdas]
    estilo = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{AZUL}")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), tamano),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0EFEC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
    )
    columnas = list(tabla.columns)
    for c, columna in enumerate(columnas):
        if pd.api.types.is_numeric_dtype(tabla[columna]):
            estilo.add("ALIGN", (c, 1), (c, -1), "RIGHT")
    fecha_venc = columnas.index("Fecha Vencimiento") if "Fecha Vencimiento" in columnas else None
    for f, fila in enumerate(celdas, start=1):
        for c, texto in enumerate(fila):
            if texto.startswith("-") and columnas[c] not in _ENTEROS:
                estilo.add("TEXTCOLOR", (c, f), (c, f), colors.HexColor(f"#{ROJO}"))
            if texto == LLAMADO:
                estilo.add("TEXTCOLOR", (c, f), (c, f), colors.HexColor(f"#{ROJO}"))
                estilo.add("FONTNAME", (c, f), (c, f), "Helvetica-Bold")
        subtotal = fecha_venc is not None and fila[fecha_venc] == "" and tabla.columns[0] == "Tipo de Forward"
        if fila[0] == "Total general" or subtotal:
            estilo.add("FONTNAME", (0, f), (-1, f), "Helvetica-Bold")
        if fila[0] == "Total general":
            estilo.add("LINEABOVE", (0, f), (-1, f), 0.8, colors.black)
    tabla_pdf = LongTable(datos, colWidths=anchos, repeatRows=1, hAlign="LEFT")
    tabla_pdf.setStyle(estilo)
    return tabla_pdf


def pdf(datos: pd.DataFrame, cuadros: list[tuple[str, pd.DataFrame]], tabla_lineas: pd.DataFrame, fecha: date) -> bytes:
    """!
    @brief PDF tamaño tabloide horizontal con las secciones Resumen, Basilea y Líneas.

    Cada sección empieza en una página nueva. En el Resumen los cuadros van uno bajo otro, sin
    partirse entre páginas; el detalle Basilea lleva todas sus columnas.
    @param datos Resultado de @ref comun.reporte_basilea.detalle.
    @param cuadros Resultado de @ref comun.reporte_basilea.resumen.
    @param tabla_lineas Resultado de @ref comun.reporte_basilea.lineas.
    @param fecha Fecha de proceso.
    @return Contenido del archivo .pdf.
    """
    salida = io.BytesIO()
    pagina = landscape(TABLOID)
    margen = 12 * mm
    ancho = pagina[0] - 2 * margen
    documento = SimpleDocTemplate(
        salida, pagesize=pagina, leftMargin=margen, rightMargin=margen, topMargin=margen, bottomMargin=margen,
        title="Reporte FX Forward Basilea", author="Gestor de Cartera de Derivados",
    )  # fmt: skip
    titulo = ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=16, leading=20)
    seccion = ParagraphStyle("seccion", fontName="Helvetica-Bold", fontSize=12, leading=16, spaceBefore=2, spaceAfter=4)
    cuadro = ParagraphStyle("cuadro", fontName="Helvetica-Bold", fontSize=9, leading=12, spaceAfter=3)
    nota = ParagraphStyle("nota", fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#52514E"))

    contenido = [
        Paragraph("Reporte FX Forward Basilea", titulo),
        Paragraph(
            escape(
                f"FX Forward vigentes al {fecha:%d-%m-%Y}. Valor razonable en CLP, desde la institución: Compra si "
                "recibe la moneda extranjera o UF y Venta si la entrega. Generado el "
                f"{datetime.now():%d-%m-%Y %H:%M}."
            ),
            nota,
        ),
        Spacer(1, 5 * mm),
        Paragraph("Resumen", seccion),
    ]
    for nombre, tabla in cuadros:
        ancho_cuadro = min(ancho, 32 * mm * len(tabla.columns))
        contenido.append(KeepTogether([Paragraph(escape(nombre), cuadro), _tabla_pdf(tabla, ancho_cuadro, 7.5), Spacer(1, 5 * mm)]))

    contenido += [PageBreak(), Paragraph("Detalle Basilea", seccion), _tabla_pdf(datos.drop(columns=["Error"]), ancho, 6.5)]
    contenido += [
        PageBreak(),
        Paragraph("Líneas por contraparte", seccion),
        Paragraph(
            escape(
                "Exposición total y situación consideran todas las operaciones vigentes de la contraparte; el valor "
                "razonable y las posiciones netas (compras menos ventas en moneda principal), sólo sus FX Forward."
            ),
            nota,
        ),
        Spacer(1, 3 * mm),
        _tabla_pdf(tabla_lineas, ancho, 7),
    ]

    def pie(lienzo, doc):
        """!
        @brief Pie de página con el nombre del reporte, la fecha de proceso y el número de página.
        @param lienzo Lienzo de reportlab.
        @param doc Documento.
        """
        lienzo.saveState()
        lienzo.setFont("Helvetica", 7)
        lienzo.setFillColor(colors.HexColor("#52514E"))
        lienzo.drawString(margen, 6 * mm, f"Reporte FX Forward Basilea · {fecha:%d-%m-%Y}")
        lienzo.drawRightString(pagina[0] - margen, 6 * mm, f"Página {doc.page}")
        lienzo.restoreState()

    documento.build(contenido, onFirstPage=pie, onLaterPages=pie)
    return salida.getvalue()

"""!
@package comun.plantilla
@brief Plantilla de carga y conversión de un Excel o CSV (una fila por pata) a operaciones de la API.
"""

import io
import re
from datetime import date, datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from comun.formato import COLUMNAS_OPERACION, COLUMNAS_PATA

## Todas las columnas de la plantilla.
COLUMNAS = COLUMNAS_OPERACION + COLUMNAS_PATA

## Columnas que se leen como enteros.
ENTEROS = {"numero", "lag_pago", "decimales_tasa"}
## Columnas que se leen como decimales; aceptan coma decimal.
DECIMALES = {"nocional", "tasa", "spread", "gearing"}
## Columnas de fecha; aceptan AAAA-MM-DD, DD-MM-AAAA y DD/MM/AAAA.
FECHAS = {"fecha_negociacion", "fecha_pago", "fecha_inicio", "fecha_fin"}
## Columnas de sí o no.
LOGICOS = {"intercambia_nocional"}
## Columnas de código que se pasan a mayúsculas.
MAYUSCULAS = {
    "producto", "modalidad_liquidacion", "tipo", "rec_pay", "moneda", "stub", "ajuste_fecha", "tipo_amortizacion",
    "fraccion_anio", "factor_capitalizacion", "periodicidad",
}  # fmt: skip
## Textos que se leen como verdadero.
VERDADERO = {"TRUE", "VERDADERO", "SI", "SÍ", "S", "1", "X"}
## Textos que se leen como falso.
FALSO = {"FALSE", "FALSO", "NO", "N", "0"}

## Columna, si es obligatoria, valores admitidos y descripción. Es la hoja Instrucciones de la plantilla.
DESCRIPCION = [
    ("deal_number", "Sí", "Letras, números, _ . -", "Folio de la operación. Se repite en sus dos filas."),
    ("producto", "Sí", "FX_FORWARD, IR_SWAP, CC_SWAP", "IR_SWAP: patas en la misma moneda. CC_SWAP: en monedas distintas."),
    ("rut_contraparte", "No", "RUT, con o sin puntos", "Debe existir en el maestro de contrapartes (página Contrapartes)."),
    ("modalidad_liquidacion", "No", "COMPENSACION, ENTREGA_FISICA", "Si se omite, COMPENSACION."),
    ("grupo_calce", "No", "Texto", "Operaciones con el mismo grupo se cubren entre sí. Vacío: descalzada."),
    ("fecha_negociacion", "Sí", "Fecha", "Fecha de cierre de la operación."),
    ("numero", "Sí", "1 o 2", "Número de la pata."),
    ("tipo", "Sí", "SIMPLE, FIJA, INDICE_OVERNIGHT", "SIMPLE es la pata de un forward."),
    ("rec_pay", "Sí", "RECIBE, PAGA", "Una pata recibe y la otra paga."),
    ("moneda", "Sí", "CLP, CLF, USD, EUR, …", "La UF va como CLF."),
    ("nocional", "Sí", "Número positivo", "Monto en la moneda de la pata."),
    ("curva_descuento", "Sí", "Código de curva", "Por ejemplo Curva CLP."),
    ("fecha_pago", "Sólo SIMPLE", "Fecha", "Fecha de liquidación del forward; igual en ambas patas."),
    ("fecha_inicio", "Patas con cupones", "Fecha", "Inicio del devengo."),
    ("fecha_fin", "Patas con cupones", "Fecha", "Vencimiento de la pata."),
    ("periodicidad", "Patas con cupones", "1M, 3M, 6M, 1Y, …", "Frecuencia de pago de cupones."),
    ("stub", "No", "NO, SHORTFRONT, SHORTBACK, LONGFRONT, LONGBACK", "Por defecto NO."),
    ("ajuste_fecha", "No", "NO, FOLLOW, MODFOLLOW, PREVIOUS, MODPREVIOUS", "Por defecto MODFOLLOW."),
    ("calendario", "No", "Código de calendario", "Por defecto SCL."),
    ("lag_pago", "No", "Entero ≥ 0", "Días hábiles entre fin de cupón y pago. Por defecto 0."),
    ("tipo_amortizacion", "No", "BULLET", "Todo el nocional al vencimiento. Por defecto BULLET."),
    ("intercambia_nocional", "No", "SI, NO", "Si el nocional se paga al vencimiento. Por defecto NO."),
    ("fraccion_anio", "No", "ACT360, ACT365, ACTACT, 30360, ACT30, 3030", "Por defecto ACT360."),
    ("factor_capitalizacion", "No", "LIN, COM, CON", "Por defecto LIN."),
    ("tasa", "Sólo FIJA", "Tanto por uno", "0.05 es 5%."),
    ("indice", "Sólo INDICE_OVERNIGHT", "ICP, SOFRINDX, …", "Índice overnight capitalizado."),
    ("curva_proyeccion", "Sólo INDICE_OVERNIGHT", "Código de curva", "Curva para proyectar el índice."),
    ("spread", "No", "Tanto por uno", "Por defecto 0."),
    ("gearing", "No", "Número", "Por defecto 1."),
    ("decimales_tasa", "No", "Entero de 0 a 12", "Decimales de la tasa equivalente. Por defecto 4."),
]

## Filas de ejemplo de la plantilla: un forward y un cross currency swap.
EJEMPLO = [
    {"deal_number": "FWD-EJEMPLO", "producto": "FX_FORWARD", "rut_contraparte": "76104729-9",
     "modalidad_liquidacion": "ENTREGA_FISICA", "fecha_negociacion": date(2026, 8, 3), "numero": 1, "tipo": "SIMPLE",
     "rec_pay": "RECIBE", "moneda": "USD", "nocional": 1_000_000, "curva_descuento": "Curva USD",
     "fecha_pago": date(2027, 2, 3)},
    {"deal_number": "FWD-EJEMPLO", "producto": "FX_FORWARD", "rut_contraparte": "76104729-9",
     "modalidad_liquidacion": "ENTREGA_FISICA", "fecha_negociacion": date(2026, 8, 3), "numero": 2, "tipo": "SIMPLE",
     "rec_pay": "PAGA", "moneda": "CLP", "nocional": 925_000_000, "curva_descuento": "Curva CLP",
     "fecha_pago": date(2027, 2, 3)},
    {"deal_number": "CCS-EJEMPLO", "producto": "CC_SWAP", "rut_contraparte": "76104729-9",
     "modalidad_liquidacion": "COMPENSACION", "fecha_negociacion": date(2026, 8, 3), "numero": 1, "tipo": "FIJA",
     "rec_pay": "RECIBE", "moneda": "CLF", "nocional": 100_000, "curva_descuento": "Curva CLF",
     "fecha_inicio": date(2026, 8, 5), "fecha_fin": date(2029, 8, 5), "periodicidad": "6M", "calendario": "SCL",
     "intercambia_nocional": "SI", "fraccion_anio": "ACT360", "factor_capitalizacion": "LIN", "tasa": 0.025},
    {"deal_number": "CCS-EJEMPLO", "producto": "CC_SWAP", "rut_contraparte": "76104729-9",
     "modalidad_liquidacion": "COMPENSACION", "fecha_negociacion": date(2026, 8, 3), "numero": 2, "tipo": "FIJA",
     "rec_pay": "PAGA", "moneda": "CLP", "nocional": 4_085_000_000, "curva_descuento": "Curva CLP",
     "fecha_inicio": date(2026, 8, 5), "fecha_fin": date(2029, 8, 5), "periodicidad": "6M", "calendario": "SCL",
     "intercambia_nocional": "SI", "fraccion_anio": "ACT360", "factor_capitalizacion": "LIN", "tasa": 0.048},
]  # fmt: skip


def _dar_formato(hoja, anchos: dict[str, int] | None = None) -> None:
    """!
    @brief Encabezado en azul, primera fila fija y ancho de columnas según el contenido.
    @param hoja Hoja de openpyxl.
    @param anchos Ancho fijo de algunas columnas, por título.
    """
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1C5CAB")
        celda.alignment = Alignment(vertical="center")
    hoja.freeze_panes = "A2"
    for i, columna in enumerate(hoja.iter_cols(min_row=1, max_row=min(hoja.max_row, 200)), start=1):
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        hoja.column_dimensions[get_column_letter(i)].width = (anchos or {}).get(columna[0].value, min(max(largo + 2, 10), 60))


def a_excel(datos: pd.DataFrame, nombre_hoja: str = "Operaciones") -> bytes:
    """!
    @brief Excel en el formato de carga, con las fechas como fechas.
    @param datos Una fila por pata.
    @param nombre_hoja Nombre de la hoja.
    @return Contenido del archivo .xlsx.
    """
    salida = io.BytesIO()
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        datos.to_excel(escritor, sheet_name=nombre_hoja, index=False)
        hoja = escritor.sheets[nombre_hoja]
        for col_idx, nombre in enumerate(datos.columns, start=1):
            if nombre in FECHAS:
                for fila in range(2, len(datos) + 2):
                    hoja.cell(fila, col_idx).number_format = "dd-mm-yyyy"
        _dar_formato(hoja)
    return salida.getvalue()


def plantilla_excel() -> bytes:
    """!
    @brief Plantilla de carga con dos operaciones de ejemplo y la hoja de instrucciones.
    @return Contenido del archivo .xlsx.
    """
    ejemplo = pd.DataFrame(EJEMPLO, columns=COLUMNAS)
    salida = io.BytesIO()
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        ejemplo.to_excel(escritor, sheet_name="Operaciones", index=False)
        hoja = escritor.sheets["Operaciones"]
        for col_idx, nombre in enumerate(COLUMNAS, start=1):
            if nombre in FECHAS:
                for fila in range(2, len(ejemplo) + 2):
                    hoja.cell(fila, col_idx).number_format = "dd-mm-yyyy"
        _dar_formato(hoja)
        pd.DataFrame(DESCRIPCION, columns=["Columna", "Obligatoria", "Valores", "Descripción"]).to_excel(
            escritor, sheet_name="Instrucciones", index=False
        )
        _dar_formato(escritor.sheets["Instrucciones"], {"Valores": 45, "Descripción": 70})
    return salida.getvalue()


def leer_archivo(nombre: str, contenido: bytes) -> pd.DataFrame:
    """!
    @brief Lee un Excel (primera hoja) o un CSV separado por coma o punto y coma.
    @param nombre Nombre del archivo, para saber su tipo.
    @param contenido Bytes del archivo.
    @return DataFrame con los nombres de columna en minúsculas.
    """
    if nombre.lower().endswith((".xlsx", ".xlsm", ".xls")):
        datos = pd.read_excel(io.BytesIO(contenido), sheet_name=0, dtype=object)
    else:
        texto = contenido.decode("utf-8-sig", errors="replace")
        separador = ";" if texto.split("\n", 1)[0].count(";") > texto.split("\n", 1)[0].count(",") else ","
        datos = pd.read_csv(io.StringIO(texto), sep=separador, dtype=str, keep_default_na=False)
    datos.columns = [str(c).strip().lower() for c in datos.columns]
    return datos.dropna(how="all")


def _vacio(valor) -> bool:
    """!
    @brief Indica si una celda está vacía.
    @param valor Valor de la celda.
    @return Verdadero si es nulo, NaN o texto en blanco.
    """
    return valor is None or (isinstance(valor, float) and pd.isna(valor)) or (isinstance(valor, str) and not valor.strip())


def _convertir(columna: str, valor):
    """!
    @brief Convierte una celda al tipo que espera la API según su columna.
    @param columna Nombre de la columna.
    @param valor Valor de la celda.
    @return Valor convertido.
    @throws ValueError Si el valor no se puede convertir.
    """
    if columna in FECHAS:
        if isinstance(valor, (datetime, pd.Timestamp)):
            return valor.date().isoformat()
        if isinstance(valor, date):
            return valor.isoformat()
        texto = str(valor).strip()[:10]
        if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}$", texto):  # 2026-08-17
            return pd.to_datetime(texto.replace("/", "-"), format="%Y-%m-%d").date().isoformat()
        return pd.to_datetime(texto, dayfirst=True).date().isoformat()  # 17-08-2026 o 17/08/2026
    if columna in ENTEROS:
        return int(float(str(valor).replace(",", ".")))
    if columna in DECIMALES:
        if isinstance(valor, (int, float)):
            return float(valor)
        texto = str(valor).strip()
        if "," in texto:  # 1.234,56 o 0,05
            texto = texto.replace(".", "").replace(",", ".")
        return float(texto)
    if columna in LOGICOS:
        if isinstance(valor, bool):
            return valor
        texto = str(valor).strip().upper()
        if texto in VERDADERO:
            return True
        if texto in FALSO:
            return False
        raise ValueError(f"valor lógico no reconocido: {valor}")
    if columna == "deal_number" and isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    return texto.upper() if columna in MAYUSCULAS else texto


def a_operaciones(datos: pd.DataFrame) -> tuple[list[dict], dict[str, str]]:
    """!
    @brief Agrupa las filas por folio y arma el cuerpo de `POST /operaciones` de cada operación.

    La validación de negocio la hace la API al recibirlas.
    @param datos Filas leídas con @ref comun.plantilla.leer_archivo.
    @return Tupla (operaciones, errores de lectura por folio).
    """
    faltantes = [c for c in ("deal_number", "producto", "fecha_negociacion", "numero", "tipo", "rec_pay") if c not in datos.columns]
    if faltantes:
        raise ValueError(f"Faltan columnas en el archivo: {', '.join(faltantes)}")
    desconocidas = [c for c in datos.columns if c not in COLUMNAS and not c.startswith("unnamed")]
    if desconocidas:
        raise ValueError(f"Columnas no reconocidas: {', '.join(desconocidas)}")

    operaciones, errores = [], {}
    datos = datos.copy()
    datos["deal_number"] = [_convertir("deal_number", v) if not _vacio(v) else "" for v in datos["deal_number"]]
    for deal, filas in datos.groupby("deal_number", sort=False):
        if not deal:
            errores["(sin deal_number)"] = f"{len(filas)} fila(s) sin deal_number"
            continue
        try:
            cabecera = {}
            for columna in COLUMNAS_OPERACION:
                valores = {_convertir(columna, v) for v in filas.get(columna, []) if not _vacio(v)}
                if len(valores) > 1:
                    raise ValueError(f"{columna} distinto entre sus filas")
                if valores:
                    cabecera[columna] = valores.pop()
            patas = []
            for _, fila in filas.iterrows():
                pata = {}
                for columna in COLUMNAS_PATA:
                    if columna in fila and not _vacio(fila[columna]):
                        try:
                            pata[columna] = _convertir(columna, fila[columna])
                        except (ValueError, TypeError) as e:
                            raise ValueError(f"pata {fila.get('numero', '?')}, {columna}: {e}") from None
                patas.append(pata)
            operaciones.append({**cabecera, "patas": sorted(patas, key=lambda p: p.get("numero", 0))})
        except (ValueError, TypeError) as e:
            errores[deal] = str(e)
    return operaciones, errores

"""!
@package comun.reportes
@brief Listado de operaciones vigentes por familia de producto, en Excel y en PDF.
"""

import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from comun.formato import MODALIDADES, PRODUCTOS, numero, par_forward, rut, tasa_o_indice

## Color de los encabezados.
AZUL = "1C5CAB"
## Columnas que se formatean como fecha.
COLUMNAS_FECHA = {"Negociación", "Fecha pago", "Inicio", "Fin", "Vencimiento"}
## Columnas que se formatean como monto, sin decimales.
COLUMNAS_MONTO = {"Nocional", "Monto recibe", "Monto paga", "MtM"}


def _mtm_columna(moneda: str | None) -> str:
    """!
    @brief Título de la columna de MtM.
    @param moneda Moneda del MtM o nada.
    @return Por ejemplo `MtM (CLP)`.
    """
    return f"MtM ({moneda})" if moneda else "MtM"


def tabla_forwards(operaciones: list[dict], mtm: dict[str, float | None], moneda_mtm: str | None) -> pd.DataFrame:
    """!
    @brief Una fila por forward.
    @param operaciones Forwards de la API.
    @param mtm MtM por folio; vacío si no se pidió.
    @param moneda_mtm Moneda del MtM.
    @return DataFrame.
    """
    filas = []
    for o in operaciones:
        patas = o["patas"]
        recibe = next(p for p in patas if p["rec_pay"] == "RECIBE")
        paga = next(p for p in patas if p["rec_pay"] == "PAGA")
        par, precio = par_forward(patas)
        fila = {
            "Folio": o["deal_number"],
            "Rut Contraparte": rut(o.get("rut_contraparte")),
            "Contraparte": o.get("contraparte") or "",
            "Modalidad": MODALIDADES.get(o["modalidad_liquidacion"], o["modalidad_liquidacion"]),
            "Negociación": pd.Timestamp(o["fecha_negociacion"]),
            "Fecha pago": pd.Timestamp(recibe["fecha_pago"]),
            "Moneda recibe": recibe["moneda"],
            "Monto recibe": recibe["nocional"],
            "Moneda paga": paga["moneda"],
            "Monto paga": paga["nocional"],
            "Par": par,
            "Precio pactado": precio,
        }
        if mtm:
            fila[_mtm_columna(moneda_mtm)] = mtm.get(o["deal_number"])
        filas.append(fila)
    return pd.DataFrame(filas)


def tabla_swaps(operaciones: list[dict], mtm: dict[str, float | None], moneda_mtm: str | None) -> pd.DataFrame:
    """!
    @brief Una fila por pata; los datos de la operación y su MtM van en cada fila.
    @param operaciones Swaps de la API.
    @param mtm MtM por folio.
    @param moneda_mtm Moneda del MtM.
    @return DataFrame.
    """
    filas = []
    for o in operaciones:
        for p in sorted(o["patas"], key=lambda p: p["numero"]):
            fila = {
                "Folio": o["deal_number"],
                "Rut Contraparte": rut(o.get("rut_contraparte")),
                "Contraparte": o.get("contraparte") or "",
                "Modalidad": MODALIDADES.get(o["modalidad_liquidacion"], o["modalidad_liquidacion"]),
                "Negociación": pd.Timestamp(o["fecha_negociacion"]),
                "Pata": p["numero"],
                "Rec/Pag": "Recibe" if p["rec_pay"] == "RECIBE" else "Paga",
                "Tipo": "Fija" if p["tipo"] == "FIJA" else "Flotante",
                "Moneda": p["moneda"],
                "Nocional": p["nocional"],
                "Inicio": pd.Timestamp(p["fecha_inicio"]),
                "Fin": pd.Timestamp(p["fecha_fin"]),
                "Periodicidad": p["periodicidad"],
                "Tasa / índice": tasa_o_indice(p),
                "Base": f"{p['fraccion_anio']} {p['factor_capitalizacion']}",
                "Interc. nocional": "Sí" if p.get("intercambia_nocional") else "No",
            }
            if mtm:
                fila[_mtm_columna(moneda_mtm)] = mtm.get(o["deal_number"])
            filas.append(fila)
    return pd.DataFrame(filas)


def tabla_familia(producto: str, operaciones: list[dict], mtm: dict, moneda_mtm: str | None) -> pd.DataFrame:
    """!
    @brief Tabla del listado de un producto.
    @param producto Código del producto.
    @param operaciones Todas las operaciones.
    @param mtm MtM por folio.
    @param moneda_mtm Moneda del MtM.
    @return DataFrame del producto.
    """
    del_producto = sorted((o for o in operaciones if o["producto"] == producto), key=lambda o: o["deal_number"])
    if producto == "FX_FORWARD":
        return tabla_forwards(del_producto, mtm, moneda_mtm)
    return tabla_swaps(del_producto, mtm, moneda_mtm)


def _subtitulo(fecha_mtm, moneda_mtm) -> str:
    """!
    @brief Texto con la fecha y moneda del MtM.
    @param fecha_mtm Fecha de proceso o nada.
    @param moneda_mtm Moneda o nada.
    @return Subtítulo del reporte.
    """
    texto = f"Operaciones vigentes. Generado el {datetime.now():%d-%m-%Y %H:%M}."
    if fecha_mtm:
        texto += f" MtM al {fecha_mtm:%d-%m-%Y} en {moneda_mtm}."
    return texto


# ------------------------------------------------------------------------------ Excel


def excel(producto: str, tabla: pd.DataFrame, fecha_mtm=None, moneda_mtm=None) -> bytes:
    """!
    @brief Excel del listado de un producto.
    @param producto Código del producto.
    @param tabla Resultado de @ref comun.reportes.tabla_familia.
    @param fecha_mtm Fecha del MtM, si se incluyó.
    @param moneda_mtm Moneda del MtM.
    @return Contenido del archivo .xlsx.
    """
    salida = io.BytesIO()
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        tabla.to_excel(escritor, sheet_name="Operaciones", index=False, startrow=3)
        hoja = escritor.sheets["Operaciones"]
        hoja["A1"] = f"{PRODUCTOS[producto]}"
        hoja["A1"].font = Font(bold=True, size=14)
        hoja["A2"] = _subtitulo(fecha_mtm, moneda_mtm)
        hoja["A2"].font = Font(italic=True, color="52514E")
        encabezado = 4
        for celda in hoja[encabezado]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=AZUL)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, nombre in enumerate(tabla.columns, start=1):
            letra = get_column_letter(i)
            formato = None
            if nombre in COLUMNAS_FECHA:
                formato, ancho = "dd-mm-yyyy", 12
            elif nombre in COLUMNAS_MONTO or nombre.startswith("MtM"):
                formato, ancho = "#,##0;[Red]-#,##0", 20
            elif nombre == "Precio pactado":
                formato, ancho = "#,##0.0000", 14
            else:
                ancho = min(max(len(nombre), *(len(str(v)) for v in tabla[nombre].head(500))) + 2, 40) if len(tabla) else 14
            hoja.column_dimensions[letra].width = ancho
            if formato:
                for fila in range(encabezado + 1, encabezado + 1 + len(tabla)):
                    hoja[f"{letra}{fila}"].number_format = formato
        hoja.freeze_panes = hoja[f"B{encabezado + 1}"]
        if len(tabla):
            hoja.auto_filter.ref = f"A{encabezado}:{get_column_letter(len(tabla.columns))}{encabezado + len(tabla)}"
    return salida.getvalue()


# -------------------------------------------------------------------------------- PDF


def _celda_pdf(nombre: str, valor) -> str:
    """!
    @brief Texto de una celda en el PDF según su columna.
    @param nombre Columna.
    @param valor Valor.
    @return Texto formateado.
    """
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    if nombre in COLUMNAS_FECHA:
        return pd.Timestamp(valor).strftime("%d-%m-%Y")
    if nombre.startswith("MtM"):
        return numero(valor, 0)
    if nombre in COLUMNAS_MONTO:
        return numero(valor, 0)
    if nombre == "Precio pactado":
        return numero(valor, 4)
    return str(valor)


def pdf(producto: str, tabla: pd.DataFrame, fecha_mtm=None, moneda_mtm=None) -> bytes:
    """!
    @brief PDF horizontal del listado de un producto.
    @param producto Código del producto.
    @param tabla Resultado de @ref comun.reportes.tabla_familia.
    @param fecha_mtm Fecha del MtM, si se incluyó.
    @param moneda_mtm Moneda del MtM.
    @return Contenido del archivo .pdf.
    """
    salida = io.BytesIO()
    pagina = landscape(A4)
    documento = SimpleDocTemplate(
        salida, pagesize=pagina, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title=f"Operaciones vigentes {PRODUCTOS[producto]}", author="Valorizador de Derivados",
    )  # fmt: skip
    titulo = ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=14, leading=18)
    subtitulo = ParagraphStyle("subtitulo", fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#52514E"))
    encabezado = ParagraphStyle("encabezado", fontName="Helvetica-Bold", fontSize=6.8, leading=8, textColor=colors.white)

    columnas = list(tabla.columns)
    datos = [[Paragraph(c, encabezado) for c in columnas]]
    datos += [[_celda_pdf(c, v) for c, v in zip(columnas, fila)] for fila in tabla.itertuples(index=False)]
    if len(datos) == 1:
        datos.append(["Sin operaciones vigentes"] + [""] * (len(columnas) - 1))

    numericas = [i for i, c in enumerate(columnas) if c in COLUMNAS_MONTO or c.startswith("MtM") or c == "Precio pactado"]
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{AZUL}")),
        ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 6.8),
        ("LEADING", (0, 1), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0EFEC")]),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor(f"#{AZUL}")),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    estilo += [("ALIGN", (i, 1), (i, -1), "RIGHT") for i in numericas]

    # Ancho natural de cada columna: el mayor entre su título y sus primeros valores, más el margen.
    ancho_util = pagina[0] - 20 * mm
    naturales = []
    for i, c in enumerate(columnas):
        titulo_ancho = max(stringWidth(palabra, "Helvetica-Bold", 6.8) for palabra in c.split())
        valores = max((stringWidth(str(f[i]), "Helvetica", 6.8) for f in datos[1:300]), default=0)
        naturales.append(max(titulo_ancho, valores) + 8)
    anchos = [a * ancho_util / sum(naturales) for a in naturales]

    tabla_pdf = LongTable(datos, colWidths=anchos, repeatRows=1)
    tabla_pdf.setStyle(TableStyle(estilo))

    def pie(lienzo, doc):
        """!
        @brief Pie de página con el nombre del reporte y el número de página.
        @param lienzo Lienzo de reportlab.
        @param doc Documento.
        """
        lienzo.saveState()
        lienzo.setFont("Helvetica", 7)
        lienzo.setFillColor(colors.HexColor("#52514E"))
        lienzo.drawString(10 * mm, 6 * mm, f"Valorizador de Derivados · {PRODUCTOS[producto]}")
        lienzo.drawRightString(pagina[0] - 10 * mm, 6 * mm, f"Página {doc.page}")
        lienzo.restoreState()

    elementos = [
        Paragraph(f"Operaciones vigentes · {PRODUCTOS[producto]}", titulo),
        Paragraph(_subtitulo(fecha_mtm, moneda_mtm) + f" {len(set(tabla['Folio'])) if len(tabla) else 0} operaciones.", subtitulo),
        Spacer(1, 4 * mm),
        tabla_pdf,
    ]
    documento.build(elementos, onFirstPage=pie, onLaterPages=pie)
    return salida.getvalue()
